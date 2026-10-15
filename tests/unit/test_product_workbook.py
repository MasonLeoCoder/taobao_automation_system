"""固定商品Excel读取、格式错误和店铺分类接入测试。"""
import tempfile
import unittest
from contextlib import ExitStack
from pathlib import Path
from unittest.mock import MagicMock, patch
from openpyxl import Workbook
from ecommerce.modules.product.workbook import load_product_workbook
from ecommerce.modules.product import service
from ecommerce.platforms.qianniu import product_shop_category as platform
from ecommerce.shared.config import load_settings
from ecommerce.platforms.qianniu.product_attributes import ApplicableSceneResult
from ecommerce.platforms.qianniu.product_shop_category import ShopCategoryResult
from ecommerce.platforms.qianniu.product_sku_media import SkuMediaResult


class ProductWorkbookTests(unittest.TestCase):
    def make_file(self, root, category="家居清洁", duplicate=False):
        file = Path(root) / "产品信息.xlsx"
        book = Workbook()
        sheet = book.active
        sheet.title = "Sheet1"
        sheet.append(["分类", "分类" if duplicate else "一口价", "商家编码", "宝贝标题", "导购标题"])
        sheet.append([category, 11, "00123", "测试商品标题", "测试导购标题"])
        sheet.append([])
        sheet.append(["销售规格", "价格", "数量", "商家编码"])
        sheet.append(["规格一", 11, 10, "T001"])
        sheet.append(["规格二", 22, 20, "T002"])
        book.save(file)
        book.close()
        return file

    def test_read_preserves_types_and_source_file(self):
        with tempfile.TemporaryDirectory() as root:
            file = self.make_file(root)
            original = file.read_bytes()
            data = load_product_workbook(file)
            self.assertEqual(data.text("分类"), "家居清洁")
            self.assertEqual(data.fields["一口价"], 11)
            self.assertEqual(data.fixed_price(), "11")
            self.assertEqual(data.fields["商家编码"], "00123")
            self.assertEqual([row["数量"] for row in data.skus], [10,20])
            self.assertEqual(data.sales_specifications(), ["规格一", "规格二"])
            self.assertEqual(file.read_bytes(), original)

    def test_missing_sheet_duplicates_formulas_and_empty_category(self):
        with tempfile.TemporaryDirectory() as root:
            with self.assertRaises(ValueError):
                load_product_workbook(self.make_file(root), "不存在")
            with self.assertRaisesRegex(ValueError, "表头重复"):
                load_product_workbook(self.make_file(root, duplicate=True))
            with self.assertRaisesRegex(ValueError, "公式"):
                load_product_workbook(self.make_file(root, category='="家居清洁"'))
            with self.assertRaisesRegex(ValueError, "非空文本"):
                load_product_workbook(self.make_file(root, category=None)).text("分类")

    def test_bad_category_prevents_browser_start(self):
        with tempfile.TemporaryDirectory() as root, patch.object(service, "browser_session") as browser:
            self.make_file(root, category=None)
            with self.assertRaises(ValueError):
                service.open_publish(load_settings(root=Path(root)), keyword="测试", category="一级 > 二级",
                                     brand="无品牌", product_info_file="产品信息.xlsx")
            browser.assert_not_called()

    def test_confirmed_category_and_title_are_filled(self):
        with tempfile.TemporaryDirectory() as root, ExitStack() as stack:
            self.make_file(root)
            page = MagicMock()
            browser = stack.enter_context(patch.object(service, "browser_session"))
            browser.return_value.__enter__.return_value = (MagicMock(),page)
            for name in ("authenticate", "click_product_nav", "click_publish_product", "select_brand_and_continue"):
                stack.enter_context(patch.object(service,name,return_value=page))
            for name in ("search_publish_products", "select_publish_category"):
                stack.enter_context(patch.object(service,name))
            selector = stack.enter_context(patch.object(service,"select_shop_categories", return_value=ShopCategoryResult(["家居清洁"], {})))
            title_writer = stack.enter_context(patch.object(service,"fill_product_title"))
            guide_writer = stack.enter_context(patch.object(service,"fill_guide_title"))
            scene_writer = stack.enter_context(patch.object(service,"select_applicable_scene",return_value=ApplicableSceneResult(None,"默认场景","Excel未提供")))
            pet_writer = stack.enter_context(patch.object(service,"select_pet_type",return_value=ApplicableSceneResult(None,"默认宠物","Excel未提供")))
            handle_writer = stack.enter_context(patch.object(service,"select_handle_shape",return_value=ApplicableSceneResult(None,"默认握柄","Excel未提供")))
            material_writer = stack.enter_context(patch.object(service,"select_adhesive_material",return_value=ApplicableSceneResult(None,"默认材质","Excel未提供")))
            strength_writer = stack.enter_context(patch.object(service,"select_adhesion_strength",return_value=ApplicableSceneResult(None,"默认粘度","Excel未提供")))
            specifications_writer = stack.enter_context(patch.object(service,"create_sales_specifications"))
            details_writer = stack.enter_context(patch.object(service,"fill_sku_details"))
            media_writer = stack.enter_context(patch.object(service,"fill_sku_media", return_value=SkuMediaResult(
                {"规格一": ["SKU搜索标题"], "规格二": ["SKU搜索标题"]},
                {"规格一": {"SKU搜索主图": "缺图", "商品规格图片": "缺图"}, "规格二": {"SKU搜索主图": "缺图", "商品规格图片": "缺图"}})))
            fixed_price_writer = stack.enter_context(patch.object(service,"fill_fixed_price"))
            merchant_writer = stack.enter_context(patch.object(service,"fill_merchant_code"))
            discount_enabler = stack.enter_context(patch.object(service,"enable_multi_item_discount"))
            warehouse_selector = stack.enter_context(patch.object(service,"select_warehouse_listing"))
            submitter = stack.enter_context(patch.object(service,"submit_with_review", return_value="宝贝发布成功"))
            ready = MagicMock()
            with self.assertNoLogs(service.__name__,level="WARNING"):
                result = service.open_publish(load_settings(root=Path(root)), keyword="测试", category="一级 > 二级",
                                              brand="无品牌", product_info_file="产品信息.xlsx", on_ready=ready)
            selector.assert_called_once_with(page,["家居清洁"])
            title_writer.assert_called_once_with(page,"测试商品标题")
            guide_writer.assert_called_once_with(page,"测试导购标题")
            scene_writer.assert_called_once_with(page,None)
            pet_writer.assert_called_once_with(page,None)
            handle_writer.assert_called_once_with(page,None)
            material_writer.assert_called_once_with(page,None)
            strength_writer.assert_called_once_with(page,None)
            specifications_writer.assert_called_once_with(page,["规格一", "规格二"])
            details_writer.assert_called_once_with(page,[
                {"销售规格": "规格一", "价格": "11", "数量": "10", "商家编码": "T001"},
                {"销售规格": "规格二", "价格": "22", "数量": "20", "商家编码": "T002"}])
            self.assertIn("SKU明细=2项",result)
            media_writer.assert_called_once_with(page,["规格一", "规格二"],unittest.mock.ANY,"测试导购标题")
            self.assertEqual(media_writer.call_args.args[2].files,{})
            self.assertIn("SKU补充信息已核对2项，跳过4项",result)
            fixed_price_writer.assert_called_once_with(page,"11")
            merchant_writer.assert_called_once_with(page,"00123")
            discount_enabler.assert_called_once_with(page)
            warehouse_selector.assert_called_once_with(page)
            self.assertIn("一口价=11元",result)
            self.assertIn("销售规格=2项",result)
            self.assertIn("商品级商家编码=00123",result)
            self.assertNotIn("货号",result)
            self.assertIn("粘性强度=默认粘度（默认第一项）",result)
            self.assertIn("粘面材质=默认材质（默认第一项）",result)
            self.assertIn("握柄形态=默认握柄（默认第一项）",result)
            self.assertIn("适用宠物类型=默认宠物（默认第一项）",result)
            self.assertIn("适用场景=默认场景（默认第一项）",result)
            page.fill.assert_not_called()
            ready.assert_called_once()
            self.assertIn("店铺分类匹配：已选1项，跳过0项",result)
            submitter.assert_not_called()
            reviewer = MagicMock()
            submitted = service.open_publish(load_settings(root=Path(root)), keyword="测试", category="一级 > 二级",
                                             brand="无品牌", product_info_file="产品信息.xlsx", submit=True, on_submit_attention=reviewer)
            submitter.assert_called_once_with(unittest.mock.ANY, page, reviewer)
            self.assertIn("商品提交已确认",submitted)
            self.assertNotIn("未提交发布",submitted)

    def test_empty_title_prevents_browser_start(self):
        from openpyxl import load_workbook
        with tempfile.TemporaryDirectory() as root, patch.object(service, "browser_session") as browser:
            file = self.make_file(root)
            book = load_workbook(file)
            book["Sheet1"]["D2"] = None
            book.save(file)
            book.close()
            with self.assertRaisesRegex(ValueError, "宝贝标题"):
                service.open_publish(load_settings(root=Path(root)), keyword="测试", category="一级 > 二级",
                                     brand="无品牌", product_info_file="产品信息.xlsx")
            browser.assert_not_called()

    def test_bad_specifications_prevent_browser_start(self):
        from openpyxl import load_workbook
        for second_name in (None, " ", "规格一"):
            with self.subTest(second_name=second_name), tempfile.TemporaryDirectory() as root, patch.object(service, "browser_session") as browser:
                file = self.make_file(root)
                book = load_workbook(file)
                book["Sheet1"]["A6"] = second_name
                book.save(file)
                book.close()
                with self.assertRaisesRegex(ValueError, "销售规格"):
                    service.open_publish(load_settings(root=Path(root)), keyword="测试", category="一级 > 二级",
                                         brand="无品牌", product_info_file="产品信息.xlsx")
                browser.assert_not_called()

    def test_sku_details_preserve_cents_zero_stock_and_text_codes(self):
        from openpyxl import load_workbook
        with tempfile.TemporaryDirectory() as root:
            file = self.make_file(root)
            book = load_workbook(file)
            sheet = book["Sheet1"]
            sheet["B5"] = "0.100"
            sheet["C5"] = 0
            sheet["D5"] = "0001"
            book.save(file)
            book.close()
            original = file.read_bytes()
            self.assertEqual(load_product_workbook(file).sku_details()[0],
                             {"销售规格": "规格一", "价格": "0.1", "数量": "0", "商家编码": "0001"})
            self.assertEqual(file.read_bytes(), original)

    def test_bad_sku_details_prevent_browser_start(self):
        from openpyxl import load_workbook
        cases = [("B5", None), ("B5", True), ("B5", "NaN"), ("B5", "-1"), ("B5", 0), ("B5", "11.001"),
                 ("C5", None), ("C5", True), ("C5", -1), ("C5", 1.5), ("D5", None), ("D5", 123)]
        for cell, value in cases:
            with self.subTest(cell=cell, value=value), tempfile.TemporaryDirectory() as root, patch.object(service, "browser_session") as browser:
                file = self.make_file(root)
                book = load_workbook(file)
                book["Sheet1"][cell] = value
                book.save(file)
                book.close()
                with self.assertRaisesRegex(ValueError, "规格一"):
                    service.open_publish(load_settings(root=Path(root)), keyword="测试", category="一级 > 二级",
                                         brand="无品牌", product_info_file="产品信息.xlsx")
                browser.assert_not_called()

    def test_bad_fixed_price_prevents_browser_start(self):
        from openpyxl import load_workbook
        for value in (None, True, "NaN", 0, -1, "11.001"):
            with self.subTest(value=value), tempfile.TemporaryDirectory() as root, patch.object(service, "browser_session") as browser:
                file = self.make_file(root)
                book = load_workbook(file)
                book["Sheet1"]["B2"] = value
                book.save(file)
                book.close()
                with self.assertRaisesRegex(ValueError, "一口价"):
                    service.open_publish(load_settings(root=Path(root)), keyword="测试", category="一级 > 二级",
                                         brand="无品牌", product_info_file="产品信息.xlsx")
                browser.assert_not_called()

    def test_empty_product_merchant_code_prevents_browser_start(self):
        from openpyxl import load_workbook
        with tempfile.TemporaryDirectory() as root, patch.object(service,"browser_session") as browser:
            file = self.make_file(root)
            book = load_workbook(file)
            book["Sheet1"]["C2"] = None
            book.save(file)
            book.close()
            with self.assertRaisesRegex(ValueError,"商家编码"):
                service.open_publish(load_settings(root=Path(root)), keyword="测试", category="一级 > 二级",
                                     brand="无品牌", product_info_file="产品信息.xlsx")
            browser.assert_not_called()

    def test_empty_guide_title_logs_skip_and_allows_other_information(self):
        from openpyxl import load_workbook
        with tempfile.TemporaryDirectory() as root, patch.object(service, "browser_session") as browser:
            file = self.make_file(root)
            book = load_workbook(file)
            book["Sheet1"]["E2"] = None
            book.save(file)
            book.close()
            browser.side_effect=RuntimeError("已到达浏览器步骤")
            with self.assertLogs(service.__name__,level="WARNING") as logs, self.assertRaisesRegex(RuntimeError,"已到达浏览器步骤"):
                service.open_publish(load_settings(root=Path(root)), keyword="测试", category="一级 > 二级",
                                     brand="无品牌", product_info_file="产品信息.xlsx")
            browser.assert_called_once()
            self.assertIn("跳过基础导购标题及每个SKU的搜索标题",logs.output[0])

