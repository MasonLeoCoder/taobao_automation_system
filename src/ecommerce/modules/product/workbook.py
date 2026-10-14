"""按固定格式读取商品Excel：首两行基础数据，销售规格表头下的规格数据。"""
import re
from dataclasses import dataclass
from decimal import Decimal, InvalidOperation
from pathlib import Path
from typing import Any
from openpyxl import load_workbook


def _price_text(value: Any, label: str) -> str:
    """校验金额并转成输入文本；参数：原值、字段说明；返回：不舍入的正金额文本。"""
    try:
        if isinstance(value, bool) or value is None:
            raise ValueError
        number = Decimal(str(value).strip())
        if not number.is_finite():
            raise ValueError
    except (InvalidOperation, ValueError):
        raise ValueError(f"{label}必须为有效数值") from None
    text = format(number, "f")
    if "." in text:
        text = text.rstrip("0").rstrip(".")
    if number <= 0 or len(text.partition(".")[2]) > 2:
        raise ValueError(f"{label}必须大于0且最多两位小数，不自动四舍五入")
    return text


@dataclass(frozen=True)
class ProductWorkbook:
    """保留源文件、工作表和原字段，页面映射由业务流程明确指定。"""
    source: Path
    sheet_name: str
    fields: dict[str, Any]
    skus: list[dict[str, Any]]

    def text(self, field: str) -> str:
        """读取明确的文本字段；参数：列名；返回：去除首尾空白的非空文本。"""
        value = self.fields.get(field)
        if not isinstance(value, str) or not value.strip():
            raise ValueError(f"商品表格字段“{field}”必须为非空文本")
        return value.strip()

    def fixed_price(self) -> str:
        """读取一口价；参数：无；返回：已校验的精确金额文本，供覆盖填写。"""
        return _price_text(self.fields.get("一口价"), "商品表格一口价")

    def shop_categories(self) -> list[str]:
        """拆分分类单元格；支持分号、逗号、顿号或换行，去重并保留输入顺序。"""
        value = self.text("分类")
        names = list(dict.fromkeys(part.strip() for part in re.split(r"[;；,，、\r\n]+", value) if part.strip()))
        if not names:
            raise ValueError("商品表格“分类”列需要至少一个分类名称")
        return names

    def sales_specifications(self) -> list[str]:
        """读取全部规格；参数：无；返回：按规格记录顺序排列的非空、唯一名称。"""
        if not self.skus:
            raise ValueError("商品表格缺少销售规格记录")
        names = []
        for index, row in enumerate(self.skus, start=1):
            value = row.get("销售规格")
            if not isinstance(value, str) or not value.strip():
                raise ValueError(f"商品表格第{index}条销售规格必须为非空文本")
            name = value.strip()
            if name in names:
                raise ValueError(f"商品表格销售规格重复：{name}")
            names.append(name)
        return names

    def sku_details(self) -> list[dict[str, str]]:
        """读取SKU明细；参数：无；返回：规格、精确价格、整数数量和文本商家编码。"""
        names = self.sales_specifications()
        result = []
        for name, row in zip(names, self.skus):
            numbers = {}
            for field in ("价格", "数量"):
                value = row.get(field)
                try:
                    if isinstance(value, bool) or value is None:
                        raise ValueError
                    number = Decimal(str(value).strip())
                    if not number.is_finite():
                        raise ValueError
                except (InvalidOperation, ValueError):
                    raise ValueError(f"规格“{name}”的{field}必须为有效数值") from None
                if field == "价格":
                    numbers[field] = _price_text(number, f"规格“{name}”的价格")
                else:
                    if number < 0 or number != number.to_integral_value():
                        raise ValueError(f"规格“{name}”的数量必须为非负整数")
                    numbers[field] = str(int(number))
            code = row.get("商家编码")
            if not isinstance(code, str) or not code.strip():
                raise ValueError(f"规格“{name}”的商家编码必须为非空文本")
            result.append({"销售规格": name, **numbers, "商家编码": code.strip()})
        return result


def _headers(row, label):
    """校验固定表头；参数：单元格元组和区域名称；返回：列索引到字段名的映射。"""
    result = {}
    for index, cell in enumerate(row):
        if cell.value is None:
            continue
        if cell.data_type == "f" or not isinstance(cell.value, str) or not cell.value.strip():
            raise ValueError(f"{label}表头必须为文本：{cell.coordinate}")
        name = cell.value.strip()
        if name in result.values():
            raise ValueError(f"{label}表头重复：{name}")
        result[index] = name
    if not result:
        raise ValueError(f"缺少{label}表头")
    return result


def _values(row, headers):
    """读取原始值，拒绝公式和未命名列，保留数值及空值含义。"""
    result = {}
    for index, cell in enumerate(row):
        if cell.data_type == "f":
            raise ValueError(f"商品资料请使用明确值，不使用公式：{cell.coordinate}")
        if index not in headers and cell.value is not None:
            raise ValueError(f"商品资料列缺少表头：{cell.coordinate}")
        if index in headers:
            result[headers[index]] = cell.value
    return result


def load_product_workbook(path: Path, sheet_name: str = "Sheet1") -> ProductWorkbook:
    """只读商品资料；参数：xlsx路径、工作表名；返回：基础字段与规格，不填写页面。"""
    path = path.resolve()
    if not path.is_file() or path.suffix.lower() != ".xlsx":
        raise ValueError(f"商品资料必须是存在的xlsx文件：{path}")
    workbook = load_workbook(path, read_only=True, data_only=False, keep_links=False)
    try:
        if sheet_name not in workbook.sheetnames:
            raise ValueError(f"商品资料缺少工作表：{sheet_name}")
        rows = list(workbook[sheet_name].iter_rows())
        if len(rows) < 2:
            raise ValueError("商品资料需要第1行表头及第2行数据")
        fields = _values(rows[1], _headers(rows[0], "基础信息"))
        sku_headers = None
        skus = []
        for row in rows[2:]:
            if row[0].value == "销售规格":
                if sku_headers is not None:
                    raise ValueError("销售规格表头出现多次，请核对固定格式")
                sku_headers = _headers(row, "销售规格")
                continue
            if sku_headers is not None and any(cell.value is not None for cell in row):
                skus.append(_values(row, sku_headers))
        return ProductWorkbook(path, sheet_name, fields, skus)
    finally:
        workbook.close()
